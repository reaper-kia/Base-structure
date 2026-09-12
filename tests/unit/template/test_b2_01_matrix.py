from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
import yaml
from docx import Document

from src.modules.documents.domain.entities import Requisite, RequisiteStatus
from src.modules.templates.application.template_service import TemplateLoader
from src.modules.templates.infra.docx_renderer import TemplateDocxRenderer

DOC_TYPES_DIR = Path("src/modules/documents/config/doc_types")
ASSETS = Path("src/modules/templates/assets")
TEMPLATE_IDS = ["classic", "modern"]


def _doc_type_ids() -> list[str]:
    return [p.stem for p in sorted(DOC_TYPES_DIR.glob("*.yaml"))]


def _spec(doc_type: str) -> dict:
    return yaml.safe_load(
        (DOC_TYPES_DIR / f"{doc_type}.yaml").read_text(encoding="utf-8")
    )


def _marker(key: str) -> str:
    return f"{key.upper()}_MARKER_42"


def _requisites(spec: dict, *, value_mode: str = "marker") -> list[Requisite]:
    reqs = []
    for r in spec.get("requisites", []):
        if value_mode == "marker":
            value = _marker(r["key"])
        else:
            value = None
        reqs.append(
            Requisite(
                key=r["key"],
                label=r.get("label", r["key"]),
                value=value,
                status=RequisiteStatus.USER_PROVIDED
                if value_mode == "marker"
                else RequisiteStatus.MISSING,
                required=bool(r.get("required")),
            )
        )
    return reqs


def _full_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in doc.sections:
        parts += [p.text for p in section.header.paragraphs]
        parts += [p.text for p in section.footer.paragraphs]
    return "\n".join(parts)


@pytest.fixture
def renderer() -> TemplateDocxRenderer:
    return TemplateDocxRenderer(ASSETS)


@pytest.mark.parametrize("template_id", TEMPLATE_IDS)
@pytest.mark.parametrize("doc_type", _doc_type_ids())
def test_matrix_all_required_markers_present(renderer, doc_type, template_id) -> None:
    spec = _spec(doc_type)
    data = renderer.render("BODY_MARKER_42", _requisites(spec), template_id)
    text = _full_text(data)

    for r in spec.get("requisites", []):
        if r.get("required"):
            assert _marker(r["key"]) in text, (
                f"{doc_type}/{template_id}: реквизит «{r['key']}» не доехал до файла"
            )
    assert "BODY_MARKER_42" in text


@pytest.mark.parametrize("template_id", TEMPLATE_IDS)
def test_missing_required_renders_label(renderer, template_id) -> None:
    spec = _spec("memo")
    data = renderer.render(
        "текст", _requisites(spec, value_mode="missing"), template_id
    )
    text = _full_text(data)

    for r in spec["requisites"]:
        if r.get("required"):
            assert f"[{r.get('label', r['key'])}]" in text


def test_optional_absent_leaves_no_trace(renderer) -> None:
    spec = _spec("memo")
    reqs = [r for r in _requisites(spec) if r.required]
    data = renderer.render("текст", reqs, "classic")
    text = _full_text(data)

    for r in spec["requisites"]:
        if not r.get("required"):
            assert _marker(r["key"]) not in text
            assert f"[{r.get('label', r['key'])}]" not in text


def test_coverage_gap_caught_at_load(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    tpl = assets / "classic"
    tpl.mkdir(parents=True)
    rules = yaml.safe_load(
        (ASSETS / "classic" / "rules.yaml").read_text(encoding="utf-8")
    )
    rules["requisites_layout"] = [
        b for b in rules["requisites_layout"] if b.get("key") != "signature"
    ]
    (tpl / "rules.yaml").write_text(
        yaml.safe_dump(rules, allow_unicode=True), encoding="utf-8"
    )
    from docx import Document

    Document().save(str(tpl / "template.docx"))

    template = TemplateLoader(assets).list_templates()[0]
    assert template.available is False
