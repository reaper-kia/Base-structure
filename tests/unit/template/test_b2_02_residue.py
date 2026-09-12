from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from src.modules.documents.domain.entities import Requisite, RequisiteStatus
from src.modules.templates.infra.docx_renderer import TemplateDocxRenderer

ASSETS = Path("src/modules/templates/assets")

VALUES = {
    "addressee": "Генеральному директору Иванову И.И.",
    "doc_date": "12.09.2026",
    "reg_number": "№ 123-ИСХ",
    "salutation": "Уважаемый Иван Иванович!",
    "subject": "О проведении инвентаризации",
    "author": "Петрова А.А.",
    "position": "Главный бухгалтер",
    "signature": "________________",
    "executor": "Сидорова В.В., тел. 123",
}


@pytest.fixture
def renderer() -> TemplateDocxRenderer:
    return TemplateDocxRenderer(ASSETS)


def _full_requisites() -> list[Requisite]:
    return [
        Requisite(
            key=k, label=k, value=v, status=RequisiteStatus.USER_PROVIDED, required=True
        )
        for k, v in VALUES.items()
    ]


def _all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_modern_has_exactly_one_table(renderer) -> None:
    data = renderer.render("Текст документа.", _full_requisites(), "modern")
    assert len(Document(BytesIO(data)).tables) == 1


def test_classic_has_no_tables(renderer) -> None:
    data = renderer.render("Текст документа.", _full_requisites(), "classic")
    assert len(Document(BytesIO(data)).tables) == 0


def test_no_sample_placeholder_when_filled(renderer) -> None:
    data = renderer.render("Текст документа.", _full_requisites(), "modern")
    text = _all_text(Document(BytesIO(data)))
    assert "[Адресат]" not in text
    assert "[Автор]" not in text


def test_file_reopens_with_non_empty_paragraphs(renderer) -> None:
    data = renderer.render("Текст документа.", _full_requisites(), "modern")
    doc = Document(BytesIO(data))
    assert any(p.text.strip() for p in doc.paragraphs)


def test_long_text_does_not_duplicate_tables(renderer) -> None:
    long_text = ("Абзац с описанием процесса инвентаризации. " * 40 + "\n\n") * 6
    assert len(long_text) > 3000
    data = renderer.render(long_text, _full_requisites(), "modern")
    assert len(Document(BytesIO(data)).tables) == 1


def test_classic_header_carries_configured_organization(renderer, monkeypatch) -> None:
    """«Верхний колонтитул — название организации» из правил организаторов.

    Название приходит из настройки стенда ORG_NAME, а не из образца
    шаблона: подставлять чужое «ООО «Ромашка»» из примера нельзя.
    """
    from src.core import config

    monkeypatch.setattr(config.settings, "org_name", "ООО «Тестовая»")

    data = renderer.render("Текст документа.", _full_requisites(), "classic")
    header_text = "\n".join(
        p.text for p in Document(BytesIO(data)).sections[0].header.paragraphs
    )

    assert header_text.strip() == "ООО «Тестовая»"
    assert "[" not in header_text


def test_unset_organization_stays_an_explicit_placeholder(renderer, monkeypatch) -> None:
    """Незаданное название — явная пометка, а не выдуманная организация."""
    from src.core import config

    monkeypatch.setattr(config.settings, "org_name", "[Название организации]")

    data = renderer.render("Текст документа.", _full_requisites(), "classic")
    header_text = "\n".join(
        p.text for p in Document(BytesIO(data)).sections[0].header.paragraphs
    )

    assert header_text.strip() == "[Название организации]"


def test_footer_has_one_real_page_field(renderer) -> None:
    """Номер страницы ровно один.

    В образце шаблона поле PAGE лежит внутри элемента управления
    содержимым, и обход по `.paragraphs` его не удалял — рядом с новым
    номером печатался второй, оставшийся от образца.
    """
    data = renderer.render("Текст документа.", _full_requisites(), "classic")
    footer = Document(BytesIO(data)).sections[0].footer

    footer_text = "\n".join(p.text for p in footer.paragraphs)
    assert "[" not in footer_text

    # Настоящий field code, а не заглушка [page], и ровно один.
    assert footer._element.xml.count(" PAGE ") == 1
