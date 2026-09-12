"""B2-13.1: round-trip валидация — parse_docx_template → TemplateLoader → available=True.

ТЗ: реальный .docx (оба встроенных + пара «кривых» — без styles.xml,
без sectPr) → parse_docx_template → записать rules.yaml →
TemplateLoader.load() → available=True.
"""

from __future__ import annotations

import xml.etree.ElementTree as ET
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
import yaml
from docx import Document

from src.modules.templates.application.template_service import TemplateLoader
from src.modules.templates.infra.docx_parser import parse_docx_template

ASSETS = Path("src/modules/templates/assets")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_DECL = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'


def _docx_to_bytes(doc: Document) -> bytes:
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _strip_styles_from_docx(data: bytes) -> bytes:
    """Согласованно удаляет word/styles.xml: саму часть, связь на неё
    в word/_rels/document.xml.rels и Override в [Content_Types].xml.

    Без удаления связи архив содержал бы «висячую» relationship —
    это повреждённый пакет, а не «DOCX без styles.xml».
    """
    src = BytesIO(data)
    dst = BytesIO()
    with (
        zipfile.ZipFile(src) as zin,
        zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for item in zin.infolist():
            if item.filename == "word/styles.xml":
                continue
            payload = zin.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                root = ET.fromstring(payload)
                for rel in list(root):
                    if (rel.get("Target") or "").endswith("styles.xml"):
                        root.remove(rel)
                payload = XML_DECL + ET.tostring(root, encoding="unicode").encode(
                    "utf-8"
                )
            elif item.filename == "[Content_Types].xml":
                root = ET.fromstring(payload)
                for node in list(root):
                    if node.get("PartName") == "/word/styles.xml":
                        root.remove(node)
                payload = XML_DECL + ET.tostring(root, encoding="unicode").encode(
                    "utf-8"
                )
            zout.writestr(item, payload)
    return dst.getvalue()


def _strip_sectpr_from_document_xml(data: bytes) -> bytes:
    """Удаляет <w:sectPr> из word/document.xml внутри архива."""
    src = BytesIO(data)
    dst = BytesIO()
    with (
        zipfile.ZipFile(src) as zin,
        zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = ET.fromstring(payload)
                body = root.find(f"{{{W}}}body")
                if body is not None:
                    for sect in list(body.findall(f"{{{W}}}sectPr")):
                        body.remove(sect)
                payload = XML_DECL + ET.tostring(root, encoding="unicode").encode(
                    "utf-8"
                )
            zout.writestr(item, payload)
    return dst.getvalue()


@pytest.fixture
def scratch_loader(tmp_path: Path) -> tuple[TemplateLoader, Path]:
    assets = tmp_path / "assets"
    assets.mkdir()
    return TemplateLoader(assets), assets


def _round_trip(loader: TemplateLoader, assets: Path, tid: str, data: bytes) -> None:
    folder = assets / tid
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "template.docx").write_bytes(data)
    rules, warnings = parse_docx_template(data)
    (folder / "rules.yaml").write_text(
        yaml.safe_dump(rules, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )
    templates = {t.id: t for t in loader.list_templates()}
    assert tid in templates, f"шаблон {tid} не загрузился. warnings: {warnings}"
    assert (
        templates[tid].available is True
    ), f"шаблон {tid} available=False — правила не прошли валидацию. warnings: {warnings}"


@pytest.mark.parametrize("template_id", ["classic", "modern"])
def test_builtin_docx_round_trip(scratch_loader, template_id: str) -> None:
    loader, assets = scratch_loader
    data = (ASSETS / template_id / "template.docx").read_bytes()
    _round_trip(loader, assets, f"parsed-{template_id}", data)


def test_docx_without_styles_xml_round_trips(scratch_loader) -> None:
    """Кривой DOCX без styles.xml: варнинг есть, round-trip → available=True."""
    loader, assets = scratch_loader
    doc = Document()
    doc.add_paragraph("Тестовый абзац без стилей")
    stripped = _strip_styles_from_docx(_docx_to_bytes(doc))

    names = zipfile.ZipFile(BytesIO(stripped)).namelist()
    assert "word/styles.xml" not in names

    rules, warnings = parse_docx_template(stripped)
    assert any("styles.xml" in w for w in warnings), f"warnings: {warnings}"

    _round_trip(loader, assets, "no-styles", stripped)


def test_docx_without_sectpr_round_trips(scratch_loader) -> None:
    """Кривой DOCX без sectPr: варнинг есть, round-trip → available=True."""
    loader, assets = scratch_loader
    doc = Document()
    doc.add_paragraph("Абзац без секционных свойств")
    stripped = _strip_sectpr_from_document_xml(_docx_to_bytes(doc))

    rules, warnings = parse_docx_template(stripped)
    assert any("sectPr" in w or "pgMar" in w for w in warnings), f"warnings: {warnings}"

    _round_trip(loader, assets, "no-sectpr", stripped)


def test_default_layout_uses_canonical_keys() -> None:
    """13.2: DEFAULT_LAYOUT парсера на канонических ключах (contracts/api.md §3.4)."""
    from src.modules.templates.infra.docx_parser import DEFAULT_LAYOUT

    canonical = {
        "doc_date",
        "position",
        "signature",
        "executor",
        "addressee",
        "author",
        "subject",
        "reg_number",
    }
    parser_keys = {block["key"] for block in DEFAULT_LAYOUT}
    non_canonical = parser_keys - canonical - {"body"}
    assert not non_canonical, f"не-канонические ключи: {sorted(non_canonical)}"
    assert {"addressee", "subject", "body"}.issubset(parser_keys)
