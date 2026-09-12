from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from src.modules.documents.domain.entities import Requisite, RequisiteStatus
from src.modules.templates.infra.docx_renderer import TemplateDocxRenderer

ASSETS = Path("src/modules/templates/assets")
VALUES = {
    "addressee": "Генеральному директору Иванову И.И.",
    "doc_date": "12.09.2026",
    "reg_number": "№ 123-ИСХ",
    "salutation": "Уважаемый Иван Иванович!",
    "subject": "О проведении инвентаризации",
    "author": "Петрова А.А.",
    "position": "Главный бухгалтер",
    "signature": "________________",
    "executor": "Сидорова В.В., тел. 123",
}
BODY = "В связи с окончанием финансового года прошу провести инвентаризацию. " * 5


@pytest.fixture
def renderer() -> TemplateDocxRenderer:
    return TemplateDocxRenderer(ASSETS)


def _reqs() -> list[Requisite]:
    return [
        Requisite(
            key=k, label=k, value=v, status=RequisiteStatus.USER_PROVIDED, required=True
        )
        for k, v in VALUES.items()
    ]


def _all_text(doc: Document) -> str:
    """Полный текст документа: параграфы + ячейки таблиц + колонтитулы."""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in doc.sections:
        parts += [p.text for p in section.header.paragraphs]
        parts += [p.text for p in section.footer.paragraphs]
    return "\n".join(parts)


def test_two_templates_same_text_layer(renderer) -> None:
    """Главный тест B2-07: одни и те же значения реквизитов в обоих документах.

    Содержание = значения реквизитов + текст тела. Порядок параграфов,
    наличие таблиц, метки «Кому:»/«От кого:» — это оформление, и оно
    различается. Проверяем только фактические данные.
    """
    reqs = _reqs()
    data_c, _ = renderer.render_from_revision(BODY, reqs, "classic")
    data_m, _ = renderer.render_from_revision(BODY, reqs, "modern")

    text_c = _all_text(Document(BytesIO(data_c)))
    text_m = _all_text(Document(BytesIO(data_m)))

    # Каждое значение реквизита должно присутствовать в обоих документах
    for key, value in VALUES.items():
        assert value in text_c, f"в classic нет значения {key}: {value}"
        assert value in text_m, f"в modern нет значения {key}: {value}"

    # Текст тела идентичен: все абзацы BODY на месте в обоих документах
    for para in BODY.split("\n\n"):
        if para.strip():
            assert para.strip() in text_c
            assert para.strip() in text_m


def test_two_templates_different_page_params(renderer) -> None:
    """Разные параметры страницы и шрифта — доказательство, что оформление реально разное."""
    reqs = _reqs()
    data_c, _ = renderer.render_from_revision(BODY, reqs, "classic")
    data_m, _ = renderer.render_from_revision(BODY, reqs, "modern")

    doc_c = Document(BytesIO(data_c))
    doc_m = Document(BytesIO(data_m))

    # Classic: Times New Roman 14, поля 30/15/20/20
    # Modern: Arial 12, поля 25/20/15/15
    assert doc_c.sections[0].left_margin.mm > doc_m.sections[0].left_margin.mm
    body_c = next(p for p in doc_c.paragraphs if BODY[:20] in p.text).runs[0]
    body_m = next(p for p in doc_m.paragraphs if BODY[:20] in p.text).runs[0]
    assert body_c.font.name == "Times New Roman"
    assert body_m.font.name == "Arial"
    assert body_c.font.size.pt == 14
    assert body_m.font.size.pt == 12


def test_no_llm_import_in_renderer() -> None:
    """Рендерер физически не импортирует LLM-клиент — доказательство разделения слоёв."""
    import src.modules.templates.infra.docx_renderer as mod
    import inspect

    source = inspect.getsource(mod)
    for forbidden in ("ml_service", "llm", "gpt", "openai", "anthropic"):
        assert forbidden not in source, (
            f"рендерер не должен импортировать/использовать {forbidden}"
        )


def test_render_from_revision_is_idempotent(renderer) -> None:
    """Два вызова с теми же входами дают байт-в-байт одинаковый результат."""
    reqs = _reqs()
    data1, _ = renderer.render_from_revision(BODY, reqs, "classic")
    data2, _ = renderer.render_from_revision(BODY, reqs, "classic")
    assert data1 == data2


def test_render_from_revision_hits_cache(renderer) -> None:
    """Второй вызов с теми же входами достаёт данные из кэша, не пересобирая."""
    reqs = _reqs()
    # Первый вызов — кэш пустой
    data1, _ = renderer.render_from_revision(BODY, reqs, "classic")
    # Проверяем, что ключ кэша появился
    assert len(renderer._revision_cache) == 1
    # Второй вызов — байты идентичны, кэш не должен вырасти
    data2, _ = renderer.render_from_revision(BODY, reqs, "classic")
    assert data1 == data2
    assert len(renderer._revision_cache) == 1
