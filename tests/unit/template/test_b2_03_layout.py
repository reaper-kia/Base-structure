from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn

from src.modules.documents.domain.entities import Requisite, RequisiteStatus
from src.modules.templates.infra.docx_renderer import TemplateDocxRenderer

ASSETS = Path("src/modules/templates/assets")
TEMPLATE_IDS = ["classic", "modern"]
JC_BY_ALIGNMENT = {
    "justify": "both",
    "both": "both",
    "left": "left",
    "center": "center",
    "right": "right",
}

VALUES = {
    "addressee": "Генеральному директору Иванову И.И.",
    "doc_date": "12.09.2026",
    "reg_number": "№ 123-ИСХ",
    "salutation": "Уважаемый Иван Иванович!",
    "subject": "О проведении инвентаризации",
    "author": "Петрова А.А.",
    "position": "Главный бухгалтер",
    "signature": "________________",
    "executor": "Сидорова В.В., тел. 123",
}
BODY_MARKER = "MARKER_BODY_42"


@pytest.fixture
def renderer() -> TemplateDocxRenderer:
    return TemplateDocxRenderer(ASSETS)


def _reqs() -> list[Requisite]:
    return [
        Requisite(
            key=k, label=k, value=v, status=RequisiteStatus.USER_PROVIDED, required=True
        )
        for k, v in VALUES.items()
    ]


def _rules(template_id: str) -> dict:
    return yaml.safe_load(
        (ASSETS / template_id / "rules.yaml").read_text(encoding="utf-8")
    )


def _render(renderer, template_id: str, body: str = BODY_MARKER) -> Document:
    data = renderer.render(body, _reqs(), template_id)
    return Document(BytesIO(data))


def _body_para(doc: Document):
    return next(p for p in doc.paragraphs if BODY_MARKER in p.text)


@pytest.mark.parametrize("template_id", TEMPLATE_IDS)
def test_body_alignment_in_xml_matches_rules(renderer, template_id) -> None:
    """Тест читает XML и сверяет w:jc с alignment из правил (п. 3.5)."""
    doc = _render(renderer, template_id, BODY_MARKER + " слово " * 200)
    para = _body_para(doc)
    pPr = para._p.get_or_add_pPr()
    jc = pPr.find(qn("w:jc")) if pPr is not None else None
    val = jc.get(qn("w:val")) if jc is not None else None

    assert val == JC_BY_ALIGNMENT[_rules(template_id)["alignment"]]


@pytest.mark.parametrize("template_id", TEMPLATE_IDS)
def test_font_and_size_match_rules(renderer, template_id) -> None:
    doc = _render(renderer, template_id)
    run = _body_para(doc).runs[0]
    rules = _rules(template_id)

    assert run.font.name == rules["font"]["family"]
    assert run.font.size.pt == rules["font"]["size_pt"]


@pytest.mark.parametrize("template_id", TEMPLATE_IDS)
def test_page_margins_match_rules(renderer, template_id) -> None:
    sec = _render(renderer, template_id).sections[0]
    page = _rules(template_id)["page"]

    assert abs(sec.left_margin.mm - page["left_mm"]) <= 0.1
    assert abs(sec.right_margin.mm - page["right_mm"]) <= 0.1
    assert abs(sec.top_margin.mm - page["top_mm"]) <= 0.1
    assert abs(sec.bottom_margin.mm - page["bottom_mm"]) <= 0.1


def test_modern_footer_has_real_subject_and_date(renderer) -> None:
    footer = "\n".join(
        p.text for p in _render(renderer, "modern").sections[0].footer.paragraphs
    )

    assert VALUES["subject"] in footer
    assert VALUES["doc_date"] in footer
    assert "[" not in footer


@pytest.mark.parametrize("template_id", TEMPLATE_IDS)
def test_no_brackets_in_headers_footers(renderer, template_id) -> None:
    doc = _render(renderer, template_id)
    for part in (doc.sections[0].header, doc.sections[0].footer):
        text = "\n".join(p.text for p in part.paragraphs)
        assert "[" not in text and "]" not in text
