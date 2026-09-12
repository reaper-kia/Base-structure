"""Программная сборка DOCX по правилам шаблона.

Оформление определяется исключительно `rules.yaml`, а не ответом ИИ:
задание требует, чтобы вид документа задавался программно (п. 1.4), и
сценарий 5 проверяет ровно это — смена шаблона меняет вид и не меняет
ни одного слова.

Модуль не знает, откуда взялся текст и реквизиты. Он принимает готовые
значения и раскладывает их по блокам `requisites_layout`.
"""

from __future__ import annotations

import hashlib
from collections import defaultdict
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from src.core.config import settings
from src.modules.documents.domain.entities import Requisite
from src.modules.documents.domain.enums import RequisiteStatus
from src.modules.templates.application.template_service import TemplateLoader
from src.modules.templates.domain.entities import Template
from src.modules.templates.infra.docx_builder import set_font

SIGNATURE_LINE = "____________________"

ALIGNMENT_BY_POSITION = {
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "top_left": WD_ALIGN_PARAGRAPH.LEFT,
    "top_right": WD_ALIGN_PARAGRAPH.RIGHT,
    "top_center": WD_ALIGN_PARAGRAPH.CENTER,
    "bottom_center": WD_ALIGN_PARAGRAPH.CENTER,
    "bottom_left": WD_ALIGN_PARAGRAPH.LEFT,
    "bottom_right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _clear_paragraphs(element) -> None:
    """Удаляет параграфы элемента (тело документа или ячейка таблицы)."""
    for paragraph in list(element.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)


def _clear_tables(element) -> None:
    """Удаляет таблицы тела: остатки образца и вторая шапка в modern."""
    for table in list(element.tables):
        table._element.getparent().remove(table._element)


def _clear_part(part) -> None:
    """Полностью очищает колонтитул.

    Обход по `.paragraphs` недостаточен: в эталонном шаблоне номер страницы
    лежит внутри элемента управления содержимым (`w:sdt`), а его параграфы
    в этот список не попадают. В результате старое поле PAGE выживало
    очистку, и рядом с новым номером страницы печатался второй.
    """
    body = part._element
    for child in list(body):
        body.remove(child)


def _safe_format(template: str, **values) -> str:
    """Подстановка, при которой неизвестный ключ даёт пустую строку."""

    class _Blank(defaultdict):
        def __missing__(self, key):
            return ""

    return template.format_map(_Blank(str, values))


def _append_page_field(run) -> None:
    """Настоящее поле PAGE: номер страницы подставляет Word, а не мы."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._element.append(begin)
    run._element.append(instruction)
    run._element.append(end)


def _is_unfilled(requisite: Requisite) -> bool:
    """Реквизит без значения помечается в документе явно (п. 1.3 задания)."""
    return (
        not requisite.value
        or requisite.status == RequisiteStatus.MISSING
        or requisite.status == RequisiteStatus.LEFT_BLANK
    )


class TemplateDocxRenderer:
    """Собирает .docx: базовый файл шаблона плюс правила раскладки."""

    def __init__(self, assets_dir: Path | str = "src/modules/templates/assets") -> None:
        self.loader = TemplateLoader(assets_dir)
        # Кэш готовых рендеров: (отпечаток текста, реквизиты, шаблон) -> байты.
        self._revision_cache: dict[tuple, bytes] = {}

    def render(
        self,
        improved_text: str,
        requisites: list[Requisite],
        template_id: str,
        doc_type_name: str = "",
    ) -> bytes:
        data, _ = self.render_with_meta(
            improved_text, requisites, template_id, doc_type_name
        )
        return data

    def render_with_meta(
        self,
        improved_text: str,
        requisites: list[Requisite],
        template_id: str,
        doc_type_name: str = "",
    ) -> tuple[bytes, Template]:
        """Рендер плюс метаданные шаблона (нужны для X-Template-Fallback)."""
        template = self.loader.load_with_fallback(template_id)
        rules = template.rules

        document = Document(template.docx_path) if template.docx_path else Document()
        _clear_paragraphs(document)
        _clear_tables(document)

        self._apply_page_settings(document, rules)
        self._apply_headers_footers(document, rules, requisites, doc_type_name)
        self._apply_layout(document, rules, requisites, improved_text, doc_type_name)

        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return buffer.getvalue(), template

    def render_from_revision(
        self,
        improved_text: str,
        requisites: list[Requisite],
        template_id: str,
        doc_type_name: str = "",
    ) -> tuple[bytes, Template]:
        """Рендер неизменной ревизии: меняется только шаблон.

        Смена шаблона не должна стоить повторной обработки текста — это
        отдельный ключ кэша и пересборка оформления, не более.
        """
        cache_key = (
            hashlib.sha256(improved_text.encode("utf-8")).hexdigest(),
            tuple((r.key, r.value, str(r.status)) for r in requisites),
            template_id,
            doc_type_name,
        )

        if cache_key not in self._revision_cache:
            data, _ = self.render_with_meta(
                improved_text, requisites, template_id, doc_type_name
            )
            self._revision_cache[cache_key] = data

        template = self.loader.load_with_fallback(template_id)
        return self._revision_cache[cache_key], template

    # --- страница и колонтитулы -------------------------------------------

    def _apply_page_settings(self, document: Document, rules: dict) -> None:
        section = document.sections[0]
        page = rules["page"]
        section.top_margin = Mm(page["top_mm"])
        section.bottom_margin = Mm(page["bottom_mm"])
        section.left_margin = Mm(page["left_mm"])
        section.right_margin = Mm(page["right_mm"])

    def _apply_headers_footers(
        self,
        document: Document,
        rules: dict,
        requisites: list[Requisite],
        doc_type_name: str,
    ) -> None:
        header_footer = rules.get("header_footer", {})
        section = document.sections[0]

        context = {r.key: (r.value or "") for r in requisites}
        context["doc_type_name"] = doc_type_name
        # Название организации — настройка стенда, а не реквизит документа.
        # Не задана -> в колонтитуле останется явная пометка.
        context["org_name"] = settings.org_name

        # Остатки образца вычищаем всегда, даже если правила не задают текст.
        _clear_part(section.header)
        _clear_part(section.footer)

        for part, key in ((section.header, "header"), (section.footer, "footer")):
            config = header_footer.get(key, {})
            text = config.get("text", "")
            if not text:
                continue

            paragraph = part.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._write_with_page_field(
                paragraph,
                text,
                context,
                config.get("font_family", rules["font"]["family"]),
                config.get("font_size_pt", rules["font"]["size_pt"] - 3),
            )

    def _write_with_page_field(
        self,
        paragraph,
        text: str,
        context: dict,
        family: str,
        size_pt: float,
    ) -> None:
        """Собирает колонтитул; токен {page} становится полем PAGE."""
        parts = text.split("{page}")

        for index, part in enumerate(parts):
            rendered = _safe_format(part, **context)
            if rendered:
                run = paragraph.add_run(rendered)
                set_font(run, family)
                run.font.size = Pt(size_pt)

            if index < len(parts) - 1:
                field_run = paragraph.add_run()
                set_font(field_run, family)
                field_run.font.size = Pt(size_pt)
                _append_page_field(field_run)

    # --- тело документа ---------------------------------------------------

    def _paragraph(self, document: Document, position: str):
        paragraph = document.add_paragraph()
        paragraph.alignment = ALIGNMENT_BY_POSITION.get(
            position, WD_ALIGN_PARAGRAPH.LEFT
        )
        # У реквизитов красной строки нет: она относится только к тексту.
        paragraph.paragraph_format.first_line_indent = Cm(0)
        return paragraph

    def _add_run(
        self,
        paragraph,
        text: str,
        family: str,
        size_pt: float,
        *,
        bold: bool = False,
        unfilled: bool = False,
    ) -> None:
        run = paragraph.add_run(text)
        set_font(run, family)
        run.font.size = Pt(size_pt)
        if bold:
            run.bold = True
        if unfilled:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def _apply_layout(
        self,
        document: Document,
        rules: dict,
        requisites: list[Requisite],
        improved_text: str,
        doc_type_name: str,
    ) -> None:
        family = rules["font"]["family"]
        size = rules["font"]["size_pt"]
        spacing = rules.get("spacing", {})
        by_key = {r.key: r for r in requisites}

        for block in rules["requisites_layout"]:
            key = block.get("key")
            block_type = block.get("type") or block.get("layout")

            if key == "body":
                self._add_body(document, improved_text, rules, family, size, spacing)
                continue

            if block_type == "doc_title":
                self._add_doc_title(document, block, doc_type_name, family, size)
                continue

            if block_type == "table":
                self._add_table(document, block, by_key, family, size)
                continue

            if block_type == "line":
                self._add_line(document, block, by_key, family, size)
                continue

            if block_type == "signature":
                self._add_signature(document, block, by_key, family, size)
                continue

            self._add_requisite(document, block, by_key, family, size)

    def _add_doc_title(
        self,
        document: Document,
        block: dict,
        doc_type_name: str,
        family: str,
        size: float,
    ) -> None:
        """Название типа документа: «СЛУЖЕБНАЯ ЗАПИСКА» над заголовком.

        Пустое имя = тип не просит печатать название (письмо начинается
        сразу с темы — см. эталонные примеры организаторов).
        """
        if not doc_type_name:
            return

        paragraph = self._paragraph(document, block.get("position", "center"))
        self._add_run(
            paragraph,
            doc_type_name.upper(),
            family,
            size,
            bold=bool(block.get("bold", True)),
        )

    def _add_requisite(
        self,
        document: Document,
        block: dict,
        by_key: dict[str, Requisite],
        family: str,
        size: float,
    ) -> None:
        requisite = by_key.get(block["key"])

        # Необязательный реквизит без значения не оставляет следа в документе.
        if requisite is None or (not requisite.required and not requisite.value):
            return

        paragraph = self._paragraph(document, block.get("position", "left"))
        unfilled = _is_unfilled(requisite)

        self._add_run(
            paragraph,
            f"[{requisite.label}]" if unfilled else requisite.value,
            family,
            block.get("font_size_pt", size),
            bold=bool(block.get("bold", False)),
            unfilled=unfilled,
        )

    def _add_line(
        self,
        document: Document,
        block: dict,
        by_key: dict[str, Requisite],
        family: str,
        size: float,
    ) -> None:
        """Несколько реквизитов в одной строке: «Дата: 14.03.2025 Номер: 12-ДЗ»."""
        parts: list[tuple[str, Requisite]] = []

        for part in block.get("parts", []):
            requisite = by_key.get(part.get("key", ""))
            if requisite is None or (not requisite.required and not requisite.value):
                continue
            parts.append((str(part.get("label", "")), requisite))

        if not parts:
            return

        paragraph = self._paragraph(document, block.get("position", "left"))
        size_pt = block.get("font_size_pt", size)

        for index, (label, requisite) in enumerate(parts):
            prefix = "" if index == 0 else "    "
            if label:
                self._add_run(paragraph, f"{prefix}{label} ", family, size_pt)
            elif prefix:
                self._add_run(paragraph, prefix, family, size_pt)

            unfilled = _is_unfilled(requisite)
            self._add_run(
                paragraph,
                f"[{requisite.label}]" if unfilled else requisite.value,
                family,
                size_pt,
                unfilled=unfilled,
            )

    def _add_signature(
        self,
        document: Document,
        block: dict,
        by_key: dict[str, Requisite],
        family: str,
        size: float,
    ) -> None:
        """Блок подписи: должность, линия подписи, И.О. Фамилия.

        Формат взят из перечня реквизитов организаторов: «Подпись —
        должность, линия подписи, И.О. Фамилия».
        """
        keys = block.get("keys", [])
        position = block.get("position", "bottom_left")
        size_pt = block.get("font_size_pt", size)

        requisites = [
            by_key[key]
            for key in keys
            if key in by_key
            and (by_key[key].required or by_key[key].value)
        ]

        if not requisites:
            return

        for index, requisite in enumerate(requisites):
            # Линия подписи ставится перед последним элементом — фамилией.
            if index == len(requisites) - 1 and len(requisites) > 1:
                line_paragraph = self._paragraph(document, position)
                self._add_run(
                    line_paragraph, block.get("line", SIGNATURE_LINE), family, size_pt
                )

            paragraph = self._paragraph(document, position)
            unfilled = _is_unfilled(requisite)
            self._add_run(
                paragraph,
                f"[{requisite.label}]" if unfilled else requisite.value,
                family,
                size_pt,
                unfilled=unfilled,
            )

    def _add_table(
        self,
        document: Document,
        block: dict,
        by_key: dict[str, Requisite],
        family: str,
        size: float,
    ) -> None:
        """Табличная шапка «Кому / От кого» шаблона «Современный регламентный»."""
        rows = block.get("rows", [])
        if not rows:
            return

        table = document.add_table(rows=len(rows), cols=2)
        table.autofit = True

        for index, row in enumerate(rows):
            label = str(row.get("label", ""))
            requisite = by_key.get(str(row.get("value_key", "")))

            if requisite is None or (not requisite.required and not requisite.value):
                value, unfilled = "", False
            elif _is_unfilled(requisite):
                value, unfilled = f"[{requisite.label}]", True
            else:
                value, unfilled = str(requisite.value), False

            label_cell = table.cell(index, 0)
            _clear_paragraphs(label_cell)
            self._add_run(label_cell.add_paragraph(), label, family, size, bold=True)

            value_cell = table.cell(index, 1)
            _clear_paragraphs(value_cell)
            if value:
                self._add_run(
                    value_cell.add_paragraph(),
                    value,
                    family,
                    size,
                    unfilled=unfilled,
                )
            else:
                value_cell.add_paragraph()

    def _add_body(
        self,
        document: Document,
        text: str,
        rules: dict,
        family: str,
        size: float,
        spacing: dict,
    ) -> None:
        alignment = ALIGNMENT_BY_POSITION.get(
            rules.get("alignment", "justify"), WD_ALIGN_PARAGRAPH.JUSTIFY
        )

        for block in (text or "").split("\n\n"):
            if not block.strip():
                continue

            paragraph = document.add_paragraph()
            paragraph.alignment = alignment

            if "line" in spacing:
                paragraph.paragraph_format.line_spacing = spacing["line"]
            if "first_line_indent_cm" in spacing:
                paragraph.paragraph_format.first_line_indent = Cm(
                    spacing["first_line_indent_cm"]
                )
            if "space_after_pt" in spacing:
                paragraph.paragraph_format.space_after = Pt(spacing["space_after_pt"])

            self._add_run(paragraph, block.strip(), family, size)
